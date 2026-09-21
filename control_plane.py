from __future__ import annotations

import hashlib
import json
import os
import re
import subprocess
import sys
import time
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from env_loader import load_project_env

ROOT = Path(__file__).resolve().parent
RUNTIME = ROOT / ".runtime" / "control"
CONFIG_PATH = RUNTIME / "config.json"
STATE_PATH = RUNTIME / "state.json"
COMMANDS_XLSX = RUNTIME / "commands.xlsx"
HEALTH_PATH = RUNTIME / "health.json"
LATEST_RESULT_PATH = RUNTIME / "latest_result.json"
RESULTS_DIR = RUNTIME / "results"
DIAGNOSTICS_DIR = RUNTIME / "diagnostics"
KORAIL_PROTECTION_MARKER = ROOT / ".runtime" / "korail_protection_failure.json"

ALLOWED_COMMANDS = {
    "status",
    "restart",
    "test_korail",
    "verify_korail_full",
    "reload_targets",
    "tail_log",
    "deploy_update",
}
REQUEST_ID_RE = re.compile(r"^[A-Za-z0-9._:-]{1,128}$")
DEFAULT_CONFIG = {
    "remote_name": "seatwatcher-drive",
    "remote_folder": "ChatGPT/SeatWatcher Control",
    "sheet_export_file": "SeatWatcher Control.xlsx",
    "poll_seconds": 30,
    "health_upload_seconds": 60,
    "protection_marker_max_age_seconds": 900,
    "rclone_config": str(RUNTIME / "rclone.conf"),
}
PROTECTION_TEXT = (
    "안정적인 환경",
    "미허가 도구",
    "매크로 등",
    "MACRO ERROR",
)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def parse_utc(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return None


def read_json(path: Path, default: Any) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError, OSError):
        return default


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(path)


def load_config() -> dict[str, Any]:
    config = dict(DEFAULT_CONFIG)
    local = read_json(CONFIG_PATH, {})
    if isinstance(local, dict):
        config.update(local)
    config["poll_seconds"] = max(15, int(config.get("poll_seconds", 30)))
    config["health_upload_seconds"] = max(30, int(config.get("health_upload_seconds", 60)))
    config["protection_marker_max_age_seconds"] = max(
        60, int(config.get("protection_marker_max_age_seconds", 900))
    )
    return config


def load_state() -> dict[str, Any]:
    state = read_json(STATE_PATH, {})
    if not isinstance(state, dict):
        state = {}
    processed = state.get("processed_request_ids")
    if not isinstance(processed, list):
        processed = []
    state["processed_request_ids"] = processed[-500:]
    return state


def save_state(state: dict[str, Any]) -> None:
    state["processed_request_ids"] = list(state.get("processed_request_ids", []))[-500:]
    write_json(STATE_PATH, state)


def run_fixed(
    args: list[str],
    *,
    timeout: int = 60,
    cwd: Path | None = None,
) -> dict[str, Any]:
    try:
        completed = subprocess.run(
            args,
            cwd=str(cwd or ROOT),
            capture_output=True,
            text=True,
            timeout=timeout,
            check=False,
        )
        return {
            "returncode": completed.returncode,
            "stdout": completed.stdout[-30000:],
            "stderr": completed.stderr[-12000:],
        }
    except subprocess.TimeoutExpired as exc:
        return {
            "returncode": 124,
            "stdout": (exc.stdout or "")[-30000:] if isinstance(exc.stdout, str) else "",
            "stderr": "timeout",
        }


def metadata_vnics() -> list[dict[str, Any]]:
    request = urllib.request.Request(
        "http://169.254.169.254/opc/v2/vnics/",
        headers={"Authorization": "Bearer Oracle"},
    )
    with urllib.request.urlopen(request, timeout=5) as response:
        payload = json.load(response)
    return payload if isinstance(payload, list) else []


def current_public_ip() -> str | None:
    try:
        for vnic in metadata_vnics():
            if vnic.get("nicIndex") == 0 or vnic.get("privateIp") == "10.0.0.112":
                value = vnic.get("publicIp")
                return str(value) if value else None
    except Exception:
        return None
    return None


def systemd_state(service: str) -> dict[str, str]:
    active = run_fixed(["systemctl", "is-active", service], timeout=10)
    enabled = run_fixed(["systemctl", "is-enabled", service], timeout=10)
    return {
        "active": active["stdout"].strip() or active["stderr"].strip(),
        "enabled": enabled["stdout"].strip() or enabled["stderr"].strip(),
    }


def drive_remote(config: dict[str, Any], relative: str) -> str:
    remote = str(config["remote_name"]).strip()
    folder = str(config["remote_folder"]).strip().strip("/")
    if not re.fullmatch(r"[A-Za-z0-9._-]+", remote):
        raise ValueError("invalid rclone remote name")
    if ".." in Path(folder).parts or ".." in Path(relative).parts:
        raise ValueError("invalid remote path")
    return f"{remote}:{folder}/{relative}"


def rclone_config_path(config: dict[str, Any]) -> str:
    return str(Path(str(config["rclone_config"])).expanduser())


def rclone_ready(config: dict[str, Any]) -> bool:
    path = Path(rclone_config_path(config))
    return path.exists() and bool(run_fixed(["rclone", "version"], timeout=10)["returncode"] == 0)


def rclone_copyto(
    config: dict[str, Any],
    source: str,
    destination: str,
    *,
    export_xlsx: bool = False,
    timeout: int = 60,
) -> dict[str, Any]:
    args = [
        "rclone",
        "copyto",
        source,
        destination,
        "--config",
        rclone_config_path(config),
    ]
    if export_xlsx:
        args.extend(["--drive-export-formats", "xlsx"])
    return run_fixed(args, timeout=timeout)


def download_command_sheet(config: dict[str, Any]) -> dict[str, Any]:
    source = drive_remote(config, str(config["sheet_export_file"]))
    COMMANDS_XLSX.parent.mkdir(parents=True, exist_ok=True)
    return rclone_copyto(
        config,
        source,
        str(COMMANDS_XLSX),
        export_xlsx=True,
        timeout=60,
    )


def upload_drive_file(
    config: dict[str, Any],
    local_path: Path,
    remote_relative: str,
    *,
    timeout: int = 60,
) -> dict[str, Any]:
    return rclone_copyto(
        config,
        str(local_path),
        drive_remote(config, remote_relative),
        timeout=timeout,
    )


def parse_commands(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise RuntimeError("openpyxl is required for the control plane") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Commands" not in workbook.sheetnames:
            raise RuntimeError("Commands sheet is missing")
        sheet = workbook["Commands"]
        headers = [sheet.cell(1, column).value for column in range(1, 7)]
        expected = [
            "request_id",
            "requested_at_utc",
            "command",
            "args_json",
            "requested_by",
            "note",
        ]
        if headers != expected:
            raise RuntimeError(f"unexpected Commands header: {headers!r}")

        commands: list[dict[str, Any]] = []
        for values in sheet.iter_rows(min_row=2, max_col=6, values_only=True):
            request_id = str(values[0] or "").strip()
            if not request_id:
                continue
            command = str(values[2] or "").strip()
            args_raw = values[3]
            if args_raw in (None, ""):
                parsed_args: dict[str, Any] = {}
            elif isinstance(args_raw, str):
                parsed = json.loads(args_raw)
                if not isinstance(parsed, dict):
                    raise ValueError("args_json must be a JSON object")
                parsed_args = parsed
            else:
                raise ValueError("args_json must be text JSON")

            commands.append(
                {
                    "request_id": request_id,
                    "requested_at_utc": str(values[1] or "").strip(),
                    "command": command,
                    "args": parsed_args,
                    "requested_by": str(values[4] or "").strip(),
                    "note": str(values[5] or "").strip(),
                }
            )
        return commands
    finally:
        workbook.close()


def tail_text(value: str, limit: int = 16000) -> str:
    return value[-limit:]


def korail_test() -> dict[str, Any]:
    python = ROOT / ".venv" / "bin" / "python"
    result = run_fixed(
        [str(python), str(ROOT / "watcher.py"), "--rail-debug"],
        timeout=120,
        cwd=ROOT,
    )
    merged = result["stdout"] + "\n" + result["stderr"]
    direct_ok = (
        result["returncode"] == 0
        and "KORAIL 완료" in merged
        and "오류 0" in merged
        and "RAIL DEGRADED" not in merged
    )
    protection = any(text in merged for text in PROTECTION_TEXT)
    return {
        "ok": direct_ok,
        "protection_block": protection,
        "returncode": result["returncode"],
        "output": tail_text(merged),
    }


def headless_check() -> dict[str, Any]:
    control_python = ROOT / ".control-venv" / "bin" / "python"
    script = ROOT / "deploy" / "korail_headless_check.py"
    if not control_python.exists():
        return {"ok": False, "error": "control_venv_missing"}
    result = run_fixed([str(control_python), str(script)], timeout=150, cwd=ROOT)
    text = result["stdout"].strip()
    try:
        payload = json.loads(text) if text else {}
    except json.JSONDecodeError:
        payload = {"raw_stdout": tail_text(text)}
    payload["runner_returncode"] = result["returncode"]
    if result["stderr"]:
        payload["runner_stderr"] = tail_text(result["stderr"], 6000)
    return payload


def command_status() -> dict[str, Any]:
    return {
        "ok": True,
        "time_utc": utc_now(),
        "public_ip": current_public_ip(),
        "seatwatcher": systemd_state("seatwatcher"),
        "control": systemd_state("seatwatcher-control"),
    }


def git_revision_state() -> dict[str, Any]:
    """Return non-secret Git/deployment revisions for mobile observability."""
    head = run_fixed(["git", "-C", str(ROOT), "rev-parse", "HEAD"], timeout=10)
    origin = run_fixed(
        ["git", "-C", str(ROOT), "rev-parse", "origin/main"],
        timeout=10,
    )
    marker_path = ROOT / ".runtime" / "origin_main_runtime_revision.txt"
    try:
        marker = marker_path.read_text(encoding="utf-8").strip()
    except OSError:
        marker = ""
    return {
        "head": head["stdout"].strip() if head["returncode"] == 0 else "",
        "origin_main": origin["stdout"].strip() if origin["returncode"] == 0 else "",
        "runtime_revision": marker,
    }


def deploy_update() -> dict[str, Any]:
    """Fetch origin/main and restart watcher so the runtime sync bridge applies it.

    No arbitrary ref, path, or shell input is accepted from Drive. The existing
    env_loader runtime bridge atomically copies the approved watcher runtime
    files from origin/main when seatwatcher.service starts.
    """
    fetched = run_fixed(
        ["git", "-C", str(ROOT), "fetch", "--prune", "origin"],
        timeout=60,
    )
    if fetched["returncode"] != 0:
        return {
            "ok": False,
            "stage": "git_fetch",
            "returncode": fetched["returncode"],
            "stderr": tail_text(fetched["stderr"] or fetched["stdout"], 6000),
            "revision": git_revision_state(),
        }

    before_restart = git_revision_state()
    origin_revision = str(before_restart.get("origin_main") or "")
    if re.fullmatch(r"[0-9a-f]{40}", origin_revision) is None:
        return {
            "ok": False,
            "stage": "origin_revision",
            "error": "invalid_origin_main_revision",
            "revision": before_restart,
        }

    restarted = restart_watcher()
    time.sleep(2)
    after_restart = git_revision_state()
    runtime_revision = str(after_restart.get("runtime_revision") or "")
    applied = runtime_revision == origin_revision
    return {
        "ok": bool(restarted.get("ok")) and applied,
        "stage": "complete" if applied else "runtime_revision_mismatch",
        "origin_main": origin_revision,
        "runtime_revision": runtime_revision,
        "seatwatcher": restarted.get("seatwatcher"),
        "restart_returncode": restarted.get("restart_returncode"),
    }


def restart_watcher() -> dict[str, Any]:
    restart = run_fixed(
        ["sudo", "-n", "systemctl", "restart", "seatwatcher"],
        timeout=30,
    )
    time.sleep(2)
    status = command_status()
    status["restart_returncode"] = restart["returncode"]
    status["restart_stderr"] = restart["stderr"]
    status["ok"] = restart["returncode"] == 0 and status["seatwatcher"]["active"] == "active"
    return status


def tail_log(args: dict[str, Any]) -> dict[str, Any]:
    try:
        lines = int(args.get("lines", 60))
    except (TypeError, ValueError):
        lines = 60
    lines = min(200, max(1, lines))
    result = run_fixed(
        [
            "journalctl",
            "-u",
            "seatwatcher",
            "-n",
            str(lines),
            "--no-pager",
            "-o",
            "short-iso",
        ],
        timeout=20,
    )
    return {
        "ok": result["returncode"] == 0,
        "lines": lines,
        "output": tail_text(result["stdout"] + result["stderr"], 24000),
    }


def execute_command(command: str, args: dict[str, Any]) -> dict[str, Any]:
    if command not in ALLOWED_COMMANDS:
        return {"ok": False, "error": f"command_not_allowed:{command}"}
    if command == "status":
        return command_status()
    if command in {"restart", "reload_targets"}:
        return restart_watcher()
    if command == "test_korail":
        return korail_test()
    if command == "verify_korail_full":
        direct = korail_test()
        headless = headless_check()
        return {
            "ok": bool(direct.get("ok")) and bool(headless.get("ok")),
            "direct": direct,
            "headless": headless,
        }
    if command == "tail_log":
        return tail_log(args)
    if command == "deploy_update":
        return deploy_update()
    raise AssertionError(command)


def result_payload(request: dict[str, Any], result: dict[str, Any]) -> dict[str, Any]:
    return {
        "request_id": request["request_id"],
        "requested_at_utc": request.get("requested_at_utc", ""),
        "processed_at_utc": utc_now(),
        "command": request["command"],
        "requested_by": request.get("requested_by", ""),
        "result": result,
    }


def persist_and_upload_result(
    config: dict[str, Any],
    payload: dict[str, Any],
) -> dict[str, Any]:
    request_id = payload["request_id"]
    safe_name = hashlib.sha256(request_id.encode("utf-8")).hexdigest()[:20]
    result_file = RESULTS_DIR / f"{safe_name}.json"
    write_json(result_file, payload)
    write_json(LATEST_RESULT_PATH, payload)
    per_request = upload_drive_file(
        config,
        result_file,
        f"results/{safe_name}.json",
    )
    latest = upload_drive_file(config, LATEST_RESULT_PATH, "latest_result.json")
    return {
        "per_request_upload": per_request["returncode"],
        "latest_upload": latest["returncode"],
    }


def build_health(
    config: dict[str, Any],
    state: dict[str, Any],
    *,
    drive_ok: bool,
    last_error: str | None = None,
) -> dict[str, Any]:
    marker = read_json(KORAIL_PROTECTION_MARKER, {})
    return {
        "updated_at_utc": utc_now(),
        "public_ip": current_public_ip(),
        "seatwatcher": systemd_state("seatwatcher"),
        "control": systemd_state("seatwatcher-control"),
        "drive_connected": drive_ok,
        "rclone_ready": rclone_ready(config),
        "protection_recovery_mode": "fallback_backoff_retry",
        "last_protection_handled_utc": state.get("last_protection_handled_utc"),
        "last_request_id": state.get("last_request_id"),
        "last_command": state.get("last_command"),
        "last_result_ok": state.get("last_result_ok"),
        "revision": git_revision_state(),
        "korail_protection_marker": marker if isinstance(marker, dict) else {},
        "last_error": last_error,
    }


def protection_marker_recent(config: dict[str, Any]) -> dict[str, Any] | None:
    marker = read_json(KORAIL_PROTECTION_MARKER, {})
    if not isinstance(marker, dict) or not marker.get("detected_at_utc"):
        return None
    detected = parse_utc(str(marker["detected_at_utc"]))
    if detected is None:
        return None
    age = (datetime.now(timezone.utc) - detected).total_seconds()
    if age < 0 or age > int(config["protection_marker_max_age_seconds"]):
        return None
    return marker


def notify_self_heal(message: str) -> None:
    try:
        from pushover_notify import is_configured, send_message

        if not is_configured():
            return
        from alert_bundle import DEFAULT_ALERT_PAGE_URL

        link_url = os.getenv("SEATWATCHER_ALERT_PAGE_URL", "").strip() or DEFAULT_ALERT_PAGE_URL
        send_message(
            message,
            link_url=link_url,
            sound="vibrate",
            title="SeatWatcher 자동 복구",
        )
    except Exception as exc:
        print(f"WARNING CONTROL pushover: {exc}", flush=True)


def maybe_auto_self_heal(
    config: dict[str, Any],
    state: dict[str, Any],
) -> dict[str, Any] | None:
    """Record one automatic recovery action for each real KORAIL protection event.

    The marker is written only after the live KORAIL direct request returns a
    protection/steady-environment response. rail_provider has already switched
    the target to NAVER degraded fallback and started the KORAIL cooldown, so the
    control plane must not immediately probe KORAIL again.
    """
    marker = protection_marker_recent(config)
    if marker is None:
        return None

    event_id = str(marker.get("event_id") or marker.get("detected_at_utc"))
    if state.get("last_protection_event_id") == event_id:
        return None

    state["last_protection_event_id"] = event_id
    state["last_protection_handled_utc"] = utc_now()
    state["last_protection_action"] = "fallback_backoff_retry"
    save_state(state)

    notify_self_heal(
        "KORAIL 보호 제한이 감지되었습니다. 공인 IP는 자동 교체하지 않고 "
        "NAVER fallback을 유지하며 cooldown 후 KORAIL 직접 조회를 다시 시험합니다."
    )
    return {
        "ok": True,
        "action": "fallback_backoff_retry",
        "event_id": event_id,
        "direct_ok": False,
        "marker": {
            "detected_at_utc": marker.get("detected_at_utc"),
            "target_id": marker.get("target_id"),
            "error_type": marker.get("error_type"),
        },
    }


def process_pending_commands(
    config: dict[str, Any],
    state: dict[str, Any],
) -> tuple[bool, str | None]:
    downloaded = download_command_sheet(config)
    if downloaded["returncode"] != 0:
        return False, tail_text(downloaded["stderr"] or downloaded["stdout"], 3000)

    try:
        commands = parse_commands(COMMANDS_XLSX)
    except Exception as exc:
        return False, f"command_sheet_parse_failed:{exc}"

    processed = set(str(item) for item in state.get("processed_request_ids", []))
    for request in commands:
        request_id = request["request_id"]
        if request_id in processed:
            continue
        if not REQUEST_ID_RE.fullmatch(request_id):
            result = {"ok": False, "error": "invalid_request_id"}
        else:
            command = request["command"]
            try:
                result = execute_command(command, request["args"])
            except Exception as exc:
                result = {
                    "ok": False,
                    "error": type(exc).__name__,
                    "message": str(exc),
                }

        payload = result_payload(request, result)
        upload = persist_and_upload_result(config, payload)
        print(
            f"CONTROL request={request_id} command={request['command']} "
            f"ok={bool(result.get('ok'))} uploads={upload}",
            flush=True,
        )
        state["processed_request_ids"].append(request_id)
        state["last_request_id"] = request_id
        state["last_command"] = request["command"]
        state["last_result_ok"] = bool(result.get("ok"))
        save_state(state)
        processed.add(request_id)

    return True, None


def run_loop() -> int:
    load_project_env()
    RUNTIME.mkdir(parents=True, exist_ok=True)
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    DIAGNOSTICS_DIR.mkdir(parents=True, exist_ok=True)
    config = load_config()
    state = load_state()
    next_health = 0.0

    print(
        "CONTROL START "
        f"poll={config['poll_seconds']}s remote={config['remote_name']} "
        f"folder={config['remote_folder']}",
        flush=True,
    )

    while True:
        started = time.monotonic()
        last_error: str | None = None
        drive_ok = False
        try:
            if rclone_ready(config):
                drive_ok, last_error = process_pending_commands(config, state)
            else:
                last_error = "rclone_not_configured"
        except Exception as exc:
            last_error = f"{type(exc).__name__}:{exc}"
            print(f"WARNING CONTROL cycle: {last_error}", flush=True)

        try:
            self_heal = maybe_auto_self_heal(config, state)
            if self_heal is not None:
                synthetic = {
                    "request_id": "auto-self-heal-" + hashlib.sha256(
                        str(self_heal).encode("utf-8")
                    ).hexdigest()[:16],
                    "requested_at_utc": utc_now(),
                    "command": "auto_protection_recovery",
                    "requested_by": "seatwatcher-control",
                }
                payload = result_payload(synthetic, self_heal)
                if rclone_ready(config):
                    persist_and_upload_result(config, payload)
        except Exception as exc:
            print(f"WARNING CONTROL self_heal: {exc}", flush=True)

        now = time.monotonic()
        if now >= next_health:
            health = build_health(
                config,
                state,
                drive_ok=drive_ok,
                last_error=last_error,
            )
            write_json(HEALTH_PATH, health)
            if rclone_ready(config):
                uploaded = upload_drive_file(config, HEALTH_PATH, "health.json")
                if uploaded["returncode"] != 0:
                    print(
                        "WARNING CONTROL health_upload: "
                        + tail_text(uploaded["stderr"] or uploaded["stdout"], 2000),
                        flush=True,
                    )
            next_health = now + int(config["health_upload_seconds"])

        elapsed = time.monotonic() - started
        sleep_seconds = max(1.0, float(config["poll_seconds"]) - elapsed)
        try:
            time.sleep(sleep_seconds)
        except KeyboardInterrupt:
            return 0


def main() -> int:
    if len(sys.argv) > 1 and sys.argv[1] == "--self-test":
        load_project_env()
        direct = korail_test()
        headless = headless_check()
        payload = {
            "ok": bool(direct.get("ok")) and bool(headless.get("ok")),
            "status": command_status(),
            "direct": direct,
            "headless": headless,
        }
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 0 if payload["ok"] else 1
    if len(sys.argv) > 1 and sys.argv[1] == "--once":
        load_project_env()
        config = load_config()
        state = load_state()
        ready = rclone_ready(config)
        drive_ok = False
        last_error = None
        if ready:
            drive_ok, last_error = process_pending_commands(config, state)
        health = build_health(config, state, drive_ok=drive_ok, last_error=last_error)
        write_json(HEALTH_PATH, health)
        print(json.dumps(health, ensure_ascii=False, indent=2))
        return 0
    return run_loop()


if __name__ == "__main__":
    raise SystemExit(main())
